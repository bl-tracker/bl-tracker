"""
BL Tracker - Awamu ya 1.5
Web app (bilingual SW/EN) ya kuongeza/kuhariri/kufuta rekodi za BL,
kutuma muhtasari wa kina kwa Telegram, ku-import Excel/CSV, na
kuambatanisha PDF/picha ya BL. Data huhifadhiwa kwenye database ya kudumu
(Neon Postgres) kupitia DATABASE_URL.
"""
import os
import io
import csv
import html
from datetime import datetime, timezone, date, timedelta
from functools import wraps

import requests
from flask import (
    Flask, request, redirect, url_for, render_template, make_response, flash,
    send_file, Response, session,
)

import db
import tesws
from carriers import detect_shipping_line

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-me-in-production")
app.config["MAX_CONTENT_LENGTH"] = 15 * 1024 * 1024  # 15 MB upload cap
app.permanent_session_lifetime = timedelta(days=14)

# ----- Config -----
TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN", "")
TELEGRAM_CHAT_ID = os.environ.get("TELEGRAM_CHAT_ID", "")
NOTIFY_TOKEN = os.environ.get("NOTIFY_TOKEN", "")
NOTIFY_LANG = os.environ.get("NOTIFY_LANG", "sw")
RUN_SCHEDULER = os.environ.get("RUN_SCHEDULER", "0") == "1"
APP_PASSWORD = os.environ.get("APP_PASSWORD", "")       # set this to enable login
OVERSTAY_DAYS = int(os.environ.get("OVERSTAY_DAYS", "21") or "21")

ALLOWED_DOC_EXT = {".pdf", ".png", ".jpg", ".jpeg", ".webp"}

from i18n import t  # noqa: E402

from werkzeug.security import generate_password_hash, check_password_hash  # noqa: E402

db.init_db()


def _seed_admin():
    """First run: if there are no users yet but APP_PASSWORD is set, create an
    'admin' account with that password so the owner can log in and add others."""
    try:
        if db.count_users() == 0 and APP_PASSWORD:
            db.create_user("admin", generate_password_hash(APP_PASSWORD), "admin")
    except Exception:
        pass


_seed_admin()


def auth_enabled():
    if APP_PASSWORD:
        return True
    try:
        return db.count_users() > 0
    except Exception:
        return False


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if auth_enabled() and not session.get("uid"):
            return redirect(url_for("login", next=request.path))
        return f(*args, **kwargs)
    return wrapper


def admin_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not auth_enabled():
            return f(*args, **kwargs)  # open mode (no users/password set)
        if not session.get("uid"):
            return redirect(url_for("login", next=request.path))
        if session.get("role") != "admin":
            flash(t("not_allowed", current_lang()), "error")
            return redirect(url_for("bl"))
        return f(*args, **kwargs)
    return wrapper


def allowed_agents():
    """Set of company (agent) names the current user may see, or None = all.
    Admins (and open mode) see everything; a viewer with no companies assigned
    also sees everything (no restriction set)."""
    if (not auth_enabled()) or session.get("role") == "admin":
        return None
    raw = (session.get("agents") or "").strip()
    if not raw:
        return None
    return {a.strip() for a in raw.split(",") if a.strip()}


def filter_allowed(records):
    allow = allowed_agents()
    if allow is None:
        return records
    return [r for r in records if (r.get("agent") or "").strip() in allow]


def _counts_from(records):
    c = {"total": len(records), "completed": 0, "in_progress": 0,
         "problem": 0, "other": 0}
    for b in records:
        cat = b.get("category") if b.get("category") in c else "other"
        c[cat if cat in ("completed", "in_progress", "problem") else "other"] += 1
    return c


def _days_since(s):
    try:
        d = datetime.strptime(str(s)[:10], "%Y-%m-%d").date()
        return (date.today() - d).days
    except Exception:
        return None


def current_lang():
    lang = request.args.get("lang") or request.cookies.get("lang") or "sw"
    return "en" if lang == "en" else "sw"


@app.context_processor
def inject_helpers():
    lang = current_lang()
    return {
        "lang": lang,
        "t": lambda key: t(key, lang),
        "cat_label": lambda c: t(c if c in db.CATEGORIES else "other", lang),
        "days_since": _days_since,
        "overstay_days": OVERSTAY_DAYS,
        "auth_on": auth_enabled(),
        "logged_in": (not auth_enabled()) or bool(session.get("uid")),
        "is_admin": (not auth_enabled()) or session.get("role") == "admin",
        "username": session.get("uname", ""),
    }


# ----- Auth -----
@app.route("/login", methods=["GET", "POST"])
def login():
    lang = current_lang()
    if not auth_enabled():
        return redirect(url_for("bl"))
    if request.method == "POST":
        uname = request.form.get("username", "").strip()
        pw = request.form.get("password", "")
        u = db.get_user(uname)
        if u and check_password_hash(u["pw_hash"], pw):
            session["uid"] = u["id"]
            session["uname"] = u["username"]
            session["role"] = u["role"]
            session["agents"] = u.get("agents") or ""
            session.permanent = True
            return redirect(request.args.get("next") or url_for("bl"))
        flash(t("login_wrong", lang), "error")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("welcome"))


# ----- User management (admin only) -----
@app.route("/users", methods=["GET", "POST"])
@admin_required
def users_page():
    lang = current_lang()
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add":
            uname = request.form.get("username", "").strip()
            pw = request.form.get("password", "")
            role = request.form.get("role", "viewer")
            agents = ",".join(a.strip() for a in request.form.getlist("agents") if a.strip())
            if not uname or not pw:
                flash(t("user_need_fields", lang), "error")
            elif db.get_user(uname):
                flash(t("user_exists", lang), "error")
            else:
                db.create_user(uname, generate_password_hash(pw), role, agents)
                flash(t("user_added", lang), "ok")
        elif action == "agents":
            uid = int(request.form.get("uid", "0") or 0)
            agents = ",".join(a.strip() for a in request.form.getlist("agents") if a.strip())
            db.set_user_agents(uid, agents)
            flash(t("user_companies_saved", lang), "ok")
        elif action == "delete":
            uid = int(request.form.get("uid", "0") or 0)
            if uid == session.get("uid"):
                flash(t("user_self_del", lang), "error")  # don't delete yourself
            else:
                db.delete_user(uid)
                flash(t("user_deleted", lang), "ok")
        elif action == "passwd":
            uid = int(request.form.get("uid", "0") or 0)
            pw = request.form.get("password", "")
            if pw:
                db.set_password(uid, generate_password_hash(pw))
                flash(t("user_pw_changed", lang), "ok")
        return redirect(url_for("users_page"))
    return render_template("users.html", users=db.list_users(),
                           all_agents=db.agents_list())


# ----- Telegram -----
def send_telegram(text):
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        return False, "missing_config"
    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage"
    try:
        resp = requests.post(
            url,
            json={"chat_id": TELEGRAM_CHAT_ID, "text": text, "parse_mode": "HTML"},
            timeout=20,
        )
        if resp.status_code == 200 and resp.json().get("ok"):
            return True, "ok"
        return False, f"{resp.status_code}: {resp.text[:200]}"
    except requests.RequestException as e:
        return False, str(e)


def send_telegram_document(data_bytes, filename, caption=""):
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        return False, "missing_config"
    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendDocument"
    try:
        resp = requests.post(
            url,
            data={"chat_id": TELEGRAM_CHAT_ID, "caption": caption[:1000], "parse_mode": "HTML"},
            files={"document": (filename, data_bytes, "application/pdf")},
            timeout=40,
        )
        if resp.status_code == 200 and resp.json().get("ok"):
            return True, "ok"
        return False, f"{resp.status_code}: {resp.text[:200]}"
    except requests.RequestException as e:
        return False, str(e)


def notify_telegram():
    """Send the grouped text summary, then attach the PDF report."""
    from report import build_report_pdf
    ok, detail = send_telegram(build_summary_text())
    try:
        # Telegram PDF = ongoing (not yet gated out) only.
        pdf = build_report_pdf(lang=NOTIFY_LANG, scope="active")
        fname = "BL_Ongoing_" + datetime.now(timezone.utc).strftime("%Y%m%d_%H%M") + ".pdf"
        ok2, detail2 = send_telegram_document(pdf, fname, caption=_pdf_caption())
    except Exception as e:  # PDF failure must not block the text summary
        ok2, detail2 = False, f"pdf_error: {e}"
    return (ok or ok2), f"text={detail}; pdf={detail2}"


def _pdf_caption():
    c = db.summary_counts()
    lang = NOTIFY_LANG
    return (f"\U0001F4C4 {t('tg_header', lang)} — {t('total', lang)}: {c['total']} | "
            f"✅ {c['completed']} | \U0001F6A2 {c['in_progress']} | "
            f"⚠️ {c['problem']} | ℹ️ {c['other']}")


CAT_EMOJI = {"completed": "✅", "in_progress": "\U0001F6A2",
             "problem": "⚠️", "other": "ℹ️"}
CAT_SORT = {"problem": 0, "in_progress": 1, "completed": 2, "other": 3}


def is_done(b):
    """A BL is 'done' ONLY once it has gated out (completion = Gate Out)."""
    return b.get("category") == "completed"


def build_summary_text(lang=None):
    """Telegram summary: ONLY active BLs (not yet gated out), grouped by agent
    (clearing company). Completed ones are excluded here — they get one final
    notice through the TANeSW sync when they gate out, then live in the app."""
    lang = lang or NOTIFY_LANG
    records = db.list_bls()
    counts = db.summary_counts()
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    no_agent = t("no_agent", lang)
    active = [b for b in records if not is_done(b)]
    by_agent = {}
    for b in active:
        key = (b.get("agent") or "").strip() or no_agent
        by_agent.setdefault(key, []).append(b)

    def agent_sort(k):
        return (k == no_agent, k.lower())

    lines = [
        f"<b>{t('tg_header', lang)}</b>", "",
        (f"{t('tg_active', lang)}: <b>{len(active)}</b>  ·  "
         f"✅ {counts['completed']}  ·  \U0001F6A2 {counts['in_progress']}  ·  "
         f"⚠️ {counts['problem']}  ·  {t('total', lang)} {counts['total']}"),
        "",
    ]
    if not active:
        lines.append(t("tg_all_done", lang))
    for agent in sorted(by_agent, key=agent_sort):
        items = sorted(by_agent[agent],
                       key=lambda b: CAT_SORT.get(b["category"], 3))
        lines.append(f"\U0001F3E2 <b>{html.escape(agent)} ({len(items)})</b>")
        for b in items:
            emoji = CAT_EMOJI.get(b["category"], "ℹ️")
            client = html.escape(b["client"] or "-")
            bl_no = html.escape(b["bl_no"] or "")
            st = f" ({html.escape(b['status'])})" if b["status"] else ""
            lines.append(f"  {emoji} {client} — {bl_no}{st}")
        lines.append("")

    lines.append(f"{t('tg_time', lang)}: {now}")
    text = "\n".join(lines)
    if len(text) > 3800:
        text = text[:3800].rsplit("\n", 1)[0] + "\n…"
    return text


# ----- Welcome / landing -----
@app.route("/")
def welcome():
    return render_template("welcome.html")


# ----- Dashboard -----
@app.route("/dashboard")
@login_required
def dashboard():
    allow = allowed_agents()
    if allow is None:
        counts = db.summary_counts()
        by_agent = db.counts_by_agent()
        by_icd = db.counts_by_icd(limit=10)
    else:
        recs = filter_allowed(db.list_bls())
        counts = _counts_from(recs)
        from collections import Counter
        ac = Counter((r.get("agent") or "").strip() or "-" for r in recs)
        ic = Counter((r.get("icd") or "").strip() for r in recs if (r.get("icd") or "").strip())
        by_agent = sorted(ac.items(), key=lambda x: (-x[1], x[0]))
        by_icd = sorted(ic.items(), key=lambda x: (-x[1], x[0]))[:10]
    cat_counts = [(c, counts[c]) for c in db.CATEGORIES]
    return render_template("dashboard.html", counts=counts,
                           by_agent=by_agent, by_icd=by_icd, cat_counts=cat_counts)


# ----- BL tracker list / summary -----
@app.route("/bl")
@login_required
def bl():
    lang = current_lang()
    search = request.args.get("q", "").strip()
    agent = request.args.get("agent", "").strip()
    allow = allowed_agents()
    # a restricted user may only filter within their own companies
    if allow is not None and agent and agent not in allow:
        agent = ""
    records = filter_allowed(db.list_bls(search or None, agent or None))
    counts = _counts_from(records)
    doc_ids = db.bls_with_documents()
    agents = db.agents_list() if allow is None else sorted(allow)

    # Group by company (agent) for the collapsible accordion. Active BLs
    # (not yet gated out) and completed ones are shown in separate sections.
    no_agent = t("no_agent", lang)

    def group(recs):
        groups = []
        for r in recs:
            key = (r.get("agent") or "").strip() or no_agent
            if not groups or groups[-1]["agent"] != key:
                groups.append({"agent": key, "items": []})
            groups[-1]["items"].append(r)
        return groups

    active_groups = group([r for r in records if not is_done(r)])
    done_groups = group([r for r in records if is_done(r)])

    return render_template("index.html",
                           active_groups=active_groups, done_groups=done_groups,
                           counts=counts, search=search, doc_ids=doc_ids,
                           agents=agents, agent=agent)


@app.route("/set-lang/<lang>")
def set_lang(lang):
    resp = make_response(redirect(request.referrer or url_for("welcome")))
    resp.set_cookie("lang", "en" if lang == "en" else "sw", max_age=60 * 60 * 24 * 365)
    return resp


# ----- Document helpers -----
def _save_uploaded_doc(bl_id):
    f = request.files.get("document")
    if not f or not f.filename:
        return
    ext = os.path.splitext(f.filename)[1].lower()
    if ext not in ALLOWED_DOC_EXT:
        return
    data = f.read()
    if not data:
        return
    db.save_document(bl_id, f.filename, f.mimetype or "application/octet-stream", data)


# ----- Add / edit / delete -----
def _form_with_carrier():
    """Form data as a dict; if Shipping Line is blank, auto-detect it from the
    Vessel name / BL number so the user doesn't have to type it."""
    data = request.form.to_dict()
    if not (data.get("shipping_line") or "").strip():
        guess = detect_shipping_line(data.get("vessel", ""), data.get("bl_no", ""))
        if guess:
            data["shipping_line"] = guess
    return data


@app.route("/add", methods=["GET", "POST"])
@admin_required
def add():
    lang = current_lang()
    if request.method == "POST":
        if not request.form.get("bl_no", "").strip():
            flash(t("required", lang), "error")
            return render_template("form.html", bl=request.form, mode="add", agents=db.agents_list())
        bl_id = db.create_bl(_form_with_carrier())
        _save_uploaded_doc(bl_id)
        flash(t("created_ok", lang), "ok")
        return redirect(url_for("bl"))
    return render_template("form.html", bl={}, mode="add", agents=db.agents_list())


@app.route("/edit/<int:bl_id>", methods=["GET", "POST"])
@admin_required
def edit(bl_id):
    lang = current_lang()
    bl = db.get_bl(bl_id)
    if not bl:
        return redirect(url_for("bl"))
    if request.method == "POST":
        if not request.form.get("bl_no", "").strip():
            flash(t("required", lang), "error")
            return render_template("form.html", bl=request.form, mode="edit", bl_id=bl_id, agents=db.agents_list())
        db.update_bl(bl_id, _form_with_carrier())
        _save_uploaded_doc(bl_id)
        flash(t("updated_ok", lang), "ok")
        return redirect(url_for("bl"))
    has_doc = db.get_document(bl_id) is not None
    return render_template("form.html", bl=bl, mode="edit", bl_id=bl_id, has_doc=has_doc, agents=db.agents_list())


@app.route("/delete/<int:bl_id>", methods=["POST"])
@admin_required
def delete(bl_id):
    lang = current_lang()
    db.delete_bl(bl_id)
    flash(t("deleted_ok", lang), "ok")
    return redirect(url_for("bl"))


# ----- View / download a BL document -----
@app.route("/document/<int:bl_id>")
@login_required
def document(bl_id):
    doc = db.get_document(bl_id)
    if not doc:
        return redirect(url_for("bl"))
    return send_file(
        io.BytesIO(doc["data"]),
        mimetype=doc["content_type"] or "application/octet-stream",
        download_name=doc["filename"] or f"bl-{bl_id}",
        as_attachment=False,
    )


# ----- Import Excel / CSV -----
IMPORT_FIELDS = ["bl_no", "client", "agent", "vessel", "shipping_line", "status", "category", "bl_date", "notes", "tansad_no", "duty_paid"]
HEADER_ALIASES = {
    "bl": "bl_no", "bl no": "bl_no", "bl_no": "bl_no", "bl number": "bl_no", "namba ya bl": "bl_no",
    "client": "client", "mteja": "client", "customer": "client",
    "agent": "agent", "kampuni": "agent", "company": "agent", "clearing agent": "agent", "clearing company": "agent",
    "vessel": "vessel", "vessel name": "vessel", "meli": "vessel", "ship": "vessel",
    "shipping_line": "shipping_line", "shipping line": "shipping_line", "s/line": "shipping_line", "carrier": "shipping_line", "line": "shipping_line",
    "status": "status", "hali": "status",
    "category": "category", "kundi": "category",
    "bl_date": "bl_date", "date": "bl_date", "tarehe": "bl_date",
    "notes": "notes", "maelezo": "notes",
    "tansad_no": "tansad_no", "tansad": "tansad_no", "tansad no": "tansad_no", "tansad no.": "tansad_no",
    "duty_paid": "duty_paid", "duty": "duty_paid", "tax paid": "duty_paid", "ushuru": "duty_paid",
}
CATEGORY_ALIASES = {
    "completed": "completed", "imekamilika": "completed", "zimekamilika": "completed", "done": "completed",
    "in_progress": "in_progress", "in progress": "in_progress", "inaendelea": "in_progress", "zinaendelea": "in_progress",
    "problem": "problem", "tatizo": "problem", "ina tatizo": "problem", "zina tatizo": "problem",
    "other": "other", "nyingine": "other",
}


def _norm_header(h):
    return HEADER_ALIASES.get((h or "").strip().lower(), (h or "").strip().lower())


def _norm_cat(v):
    return CATEGORY_ALIASES.get((v or "").strip().lower(), "in_progress")


def _parse_rows(headers, data_rows):
    idx = {}
    for i, h in enumerate(headers):
        key = _norm_header(h)
        if key in IMPORT_FIELDS and key not in idx:
            idx[key] = i
    rows = []
    for raw in data_rows:
        if not any((str(c).strip() if c is not None else "") for c in raw):
            continue
        rec = {}
        for field, i in idx.items():
            val = raw[i] if i < len(raw) else ""
            rec[field] = "" if val is None else str(val).strip()
        # BL number is required — skip rows that don't have one yet.
        if not (rec.get("bl_no") or "").strip() or rec["bl_no"].lower() in ("-", "n/a", "none"):
            continue
        if rec.get("category"):
            rec["category"] = _norm_cat(rec["category"])
        rows.append(rec)
    return rows


@app.route("/import", methods=["GET", "POST"])
@admin_required
def import_bls():
    lang = current_lang()
    if request.method == "POST":
        f = request.files.get("file")
        if not f or not f.filename:
            flash(t("import_nofile", lang), "error")
            return render_template("import.html")
        name = f.filename.lower()
        data = f.read()
        dservices = False
        try:
            if name.endswith(".csv"):
                text_content = data.decode("utf-8-sig", errors="replace")
                reader = list(csv.reader(io.StringIO(text_content)))
                if not reader:
                    raise ValueError("empty")
                rows = _parse_rows(reader[0], reader[1:])
            elif name.endswith(".xlsx"):
                from openpyxl import load_workbook
                import dservices_import
                wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                if dservices_import.is_dservices_workbook(wb.sheetnames):
                    dservices = True
                    rows = dservices_import.parse(data)
                else:
                    ws = wb.active
                    all_rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
                    if not all_rows:
                        raise ValueError("empty")
                    rows = _parse_rows(all_rows[0], all_rows[1:])
            else:
                flash(t("import_badtype", lang), "error")
                return render_template("import.html")
        except Exception:
            flash(t("import_error", lang), "error")
            return render_template("import.html")

        if request.form.get("replace"):
            db.clear_all_bls()
        n = db.bulk_create(rows)
        flash(t("import_ds_ok" if dservices else "import_ok", lang).format(n=n), "ok")
        return redirect(url_for("bl"))
    return render_template("import.html")


@app.route("/template.csv")
@login_required
def template_csv():
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(IMPORT_FIELDS)
    w.writerow(["YSPQINMTW201", "ABC Traders Ltd", "Clearing Co A", "MAERSK NADI", "Maersk",
                "Cargo at port", "in_progress", "2026-10-10", "Inasubiri assessment"])
    w.writerow(["GOSUNGB20807248", "Juma Imports", "Clearing Co B", "CALANDRA", "MSC",
                "Released", "completed", "2026-10-05", ""])
    resp = Response(out.getvalue(), mimetype="text/csv")
    resp.headers["Content-Disposition"] = "attachment; filename=bl_template.csv"
    return resp


# ----- Send summary now (UI button) -----
@app.route("/send-summary", methods=["POST"])
@admin_required
def send_summary():
    lang = current_lang()
    ok, _ = notify_telegram()
    flash(t("summary_sent" if ok else "summary_fail", lang), "ok" if ok else "error")
    return redirect(url_for("bl"))


# ----- PDF report (download / print) -----
@app.route("/report.pdf")
@login_required
def report_pdf():
    from report import build_report_pdf
    agent = request.args.get("agent", "").strip() or None
    scope = request.args.get("scope", "").strip() or None
    if scope not in ("active", "done"):
        scope = None
    allow = allowed_agents()
    if allow is not None and agent and agent not in allow:
        agent = None
    lang = current_lang()
    pdf = build_report_pdf(agent_filter=agent, lang=lang, scope=scope, allowed=allow)
    resp = make_response(pdf)
    resp.headers["Content-Type"] = "application/pdf"
    tag = {"active": "Ongoing", "done": "Completed"}.get(scope, "All")
    fname = f"BL_{tag}_" + datetime.now(timezone.utc).strftime("%Y%m%d") + ".pdf"
    resp.headers["Content-Disposition"] = f'inline; filename="{fname}"'
    return resp


# ----- TANeSW auto-sync -----
_FIELD_LABELS = {
    "status": "Status", "vessel": "Vessel", "shipping_line": "S/Line",
    "icd": "ICD", "manifest_date": "Manifest", "carry_in_date": "Carry-in",
    "release_date": "Release", "gate_out_date": "Gate-out",
    "tansad_no": "TANSAD", "containers": "Containers", "container_count": "Cntrs",
}


import threading

_sync_lock = threading.Lock()
_sync_state = {"running": False, "last": None}


def sync_tanesw(only_open=True):
    """Fetch live details from TANeSW for each BL and apply changes.
    Returns (result_dict, changes_list)."""
    rows = db.list_bls()
    changes = tesws.sync_rows(rows, only_open=only_open)
    for ch in changes:
        db.apply_tesws_update(ch["id"], ch["fields"])
    checked = sum(1 for r in rows
                  if not (only_open and r.get("category") == "completed"))
    return ({"checked": checked, "updated": len(changes), "total": len(rows)},
            changes)


def _sync_worker(only_open, do_notify):
    """Run the sync off the request thread (TANeSW has many BLs; this can take
    a minute). Sends a Telegram note when something changed, or on error."""
    try:
        result, changes = sync_tanesw(only_open=only_open)
        if do_notify and changes:
            send_telegram(build_changes_text(changes))
        _sync_state["last"] = {"ok": True, **result}
    except Exception as e:  # never die silently — report to logs + Telegram
        import traceback
        traceback.print_exc()
        _sync_state["last"] = {"ok": False, "error": str(e)}
        try:
            send_telegram(f"⚠️ TANeSW sync error: {html.escape(str(e))[:300]}")
        except Exception:
            pass
    finally:
        _sync_state["running"] = False


def start_sync(only_open=True, do_notify=True):
    """Kick off a background sync if one isn't already running.
    Returns True if started, False if one was already in progress."""
    with _sync_lock:
        if _sync_state["running"]:
            return False
        _sync_state["running"] = True
    threading.Thread(target=_sync_worker, args=(only_open, do_notify),
                     daemon=True).start()
    return True


def build_changes_text(changes, lang=None):
    lang = lang or NOTIFY_LANG
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    if not changes:
        return (f"<b>{t('tg_sync_header', lang)}</b>\n\n"
                f"{t('tg_sync_none', lang)}\n\n{t('tg_time', lang)}: {now}")
    lines = [f"<b>{t('tg_sync_header', lang)}</b>", ""]

    # Highlight newly-manifested BLs first (these just got a manifest date).
    newly = [c for c in changes if c.get("new_manifest")]
    if newly:
        lines.append(f"\U0001F195 <b>{t('tg_new_manifest', lang)}</b>")
        for ch in newly[:30]:
            client = html.escape(ch.get("client") or "-")
            bl_no = html.escape(ch.get("bl_no") or "")
            mdate = html.escape(str(ch["changes"]["manifest_date"][1]))
            lines.append(f"   \U0001F4C5 {client} — {bl_no} ({mdate})")
        lines.append("")

    for ch in changes[:40]:
        detail = []
        for field, (old, new) in ch["changes"].items():
            if field == "category" or not new:  # internal / cleared values
                continue
            lbl = _FIELD_LABELS.get(field, field)
            detail.append(f"   {html.escape(lbl)}: {html.escape(str(new))}")
        if not detail:
            continue
        client = html.escape(ch.get("client") or "-")
        bl_no = html.escape(ch.get("bl_no") or "")
        lines.append(f"\U0001F504 <b>{client}</b> — {bl_no}")
        lines.extend(detail)
        lines.append("")
    lines.append(f"{t('tg_time', lang)}: {now}")
    text = "\n".join(lines)
    if len(text) > 3800:
        text = text[:3800].rsplit("\n", 1)[0] + "\n…"
    return text


# ----- Weekly report -----
def build_weekly_text(lang=None, window=7):
    """A once-a-week recap: gated out this week, newly manifested this week,
    what's still ongoing, and overstays — grouped by company."""
    lang = lang or NOTIFY_LANG
    records = db.list_bls()
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    def within(dstr):
        d = _days_since(dstr)
        return d is not None and 0 <= d < window

    gated = [b for b in records if within(b.get("gate_out_date"))]
    manif = [b for b in records if within(b.get("manifest_date"))]
    ongoing = [b for b in records if not is_done(b)]
    overdue = [b for b in ongoing
               if (_days_since(b.get("manifest_date") or b.get("carry_in_date"))
                   or -1) >= OVERSTAY_DAYS]

    no_agent = t("no_agent", lang)

    def by_company(items):
        g = {}
        for b in items:
            g.setdefault((b.get("agent") or "").strip() or no_agent, []).append(b)
        return g

    lines = [f"<b>{t('wk_header', lang)}</b>", "",
             f"{t('wk_gated', lang)}: <b>{len(gated)}</b>  ·  "
             f"\U0001F195 {t('wk_manifested', lang)}: <b>{len(manif)}</b>",
             f"\U0001F6A2 {t('wk_ongoing', lang)}: <b>{len(ongoing)}</b>  ·  "
             f"⚠️ {t('overstay', lang)}: <b>{len(overdue)}</b>", ""]

    def section(title, items, emoji):
        if not items:
            return
        lines.append(f"{emoji} <b>{title}</b>")
        for agent in sorted(by_company(items), key=lambda k: (k == no_agent, k.lower())):
            grp = by_company(items)[agent]
            lines.append(f"\U0001F3E2 {html.escape(agent)} ({len(grp)})")
            for b in grp[:20]:
                lines.append(f"   {html.escape(b.get('client') or '-')} — "
                             f"{html.escape(b.get('bl_no') or '')}")
        lines.append("")

    section(t("wk_gated", lang), gated, "✅")
    section(t("wk_manifested", lang), manif, "\U0001F195")
    if overdue:
        section(t("overstay", lang), overdue, "⚠️")

    lines.append(f"{t('tg_time', lang)}: {now}")
    text = "\n".join(lines)
    if len(text) > 3800:
        text = text[:3800].rsplit("\n", 1)[0] + "\n…"
    return text


def weekly_report():
    ok, detail = send_telegram(build_weekly_text())
    try:
        from report import build_report_pdf
        pdf = build_report_pdf(lang=NOTIFY_LANG, scope="active")
        fname = "BL_Weekly_" + datetime.now(timezone.utc).strftime("%Y%m%d") + ".pdf"
        ok2, detail2 = send_telegram_document(pdf, fname, caption=t("wk_header", NOTIFY_LANG))
    except Exception as e:
        ok2, detail2 = False, f"pdf_error: {e}"
    return (ok or ok2), f"text={detail}; pdf={detail2}"


@app.route("/tasks/weekly")
def tasks_weekly():
    if NOTIFY_TOKEN and request.args.get("token") != NOTIFY_TOKEN:
        return {"ok": False, "error": "unauthorized"}, 401
    ok, detail = weekly_report()
    return {"ok": ok, "detail": detail}


# ----- Notify endpoint for external cron (every 3 hours) -----
@app.route("/tasks/notify")
def tasks_notify():
    if NOTIFY_TOKEN and request.args.get("token") != NOTIFY_TOKEN:
        return {"ok": False, "error": "unauthorized"}, 401
    ok, detail = notify_telegram()
    return {"ok": ok, "detail": detail}


@app.route("/summary.json")
def summary_json():
    if NOTIFY_TOKEN and request.args.get("token") != NOTIFY_TOKEN:
        return {"ok": False, "error": "unauthorized"}, 401
    return {"ok": True, "counts": db.summary_counts()}


# ----- TANeSW sync endpoint for external cron (e.g. hourly) -----
@app.route("/tasks/sync")
def tasks_sync():
    if NOTIFY_TOKEN and request.args.get("token") != NOTIFY_TOKEN:
        return {"ok": False, "error": "unauthorized"}, 401
    started = start_sync(only_open=request.args.get("all") != "1",
                         do_notify=request.args.get("notify") != "0")
    return {"ok": True, "started": started, "last": _sync_state.get("last")}


# ----- Manual sync button (from the BL page) -----
@app.route("/sync", methods=["POST"])
@admin_required
def sync_now():
    lang = current_lang()
    started = start_sync(only_open=request.form.get("all") != "1", do_notify=True)
    flash(t("sync_started" if started else "sync_running", lang), "ok")
    return redirect(url_for("bl"))


# ----- Debug: raw TANeSW mapped fields for one BL (token) -----
@app.route("/tesws/raw")
def tesws_raw():
    if NOTIFY_TOKEN and request.args.get("token") != NOTIFY_TOKEN:
        return {"ok": False, "error": "unauthorized"}, 401
    bl = request.args.get("bl", "").strip()
    if not bl:
        return {"ok": False, "error": "pass ?bl=<BL number>"}, 400
    if request.args.get("debug") == "1":
        return {"ok": True, **tesws.diagnose(bl)}
    fields = tesws.fetch_bl(bl)
    return {"ok": True, "bl_no": bl, "found": fields is not None, "fields": fields}


@app.route("/health")
def health():
    return {"ok": True}


def _start_scheduler():
    from apscheduler.schedulers.background import BackgroundScheduler
    scheduler = BackgroundScheduler(daemon=True, timezone="UTC")
    scheduler.add_job(notify_telegram, "interval", hours=3)
    scheduler.start()


if RUN_SCHEDULER:
    _start_scheduler()


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)
